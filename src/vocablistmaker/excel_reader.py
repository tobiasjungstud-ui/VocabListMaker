"""Robustes Einlesen der Excel-Wortliste inklusive zuverlässiger Unit-Erkennung.

Die Wortlisten der Lehrmittel sind selten sauber tabellarisch aufgebaut. Dieses
Modul behandelt deshalb gezielt die typischen Fehlerquellen:

* Titel- und Leerzeilen oberhalb der eigentlichen Kopfzeile
* Blocküberschriften ("Unit 4", "Culture", "Projects"), die in Spalte A stehen
  und keine Vokabel sind
* Sektionsangaben mit Seitenzahl ("Starter Unit, p.4"), deren Zahl am Ende
  *nicht* die Unit ist - die häufigste Ursache für falsche Unit-Zuordnungen
* Zusatzbereiche, die zur selben Unit gehören ("Curriculum extra 4",
  "Culture 4", "Extra Listening and Speaking Unit 4", "Project 4", "Songs 4")
* Doppelte Einträge und Zeilen ohne Übersetzung
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import IO

from openpyxl import load_workbook

from .models import VocabEntry

# ", p.4" / ", pp. 12-13" / ", S. 7" am Ende einer Sektionsangabe.
_PAGE_SUFFIX = re.compile(
    r"[,;]\s*(?:p{1,2}\.?|pages?|seiten?|s\.)\s*[\dIVXivx]+\s*[-–—]?\s*[\dIVXivx]*\s*$",
    re.IGNORECASE,
)
# Trailing "(p. 12)"
_PAGE_PAREN = re.compile(r"\(\s*(?:p{1,2}\.?|s\.)\s*[^)]*\)\s*$", re.IGNORECASE)

_STARTER = re.compile(r"^\s*starter\b", re.IGNORECASE)
_UNIT_NUM = re.compile(
    r"(?:unit|einheit|extra|culture|kultur|project|projects|song|songs|"
    r"curriculum\s+extra|lektion)\s*(\d{1,2})\s*$",
    re.IGNORECASE,
)
_TRAILING_NUM = re.compile(r"(\d{1,2})\s*$")

_HEADER_ALIASES = {
    "english": {"english", "englisch", "word", "wort", "englisches wort", "en"},
    "german": {
        "translation",
        "german",
        "deutsch",
        "übersetzung",
        "uebersetzung",
        "deutsche übersetzung",
        "de",
    },
    "section": {"section", "sektion", "unit", "abschnitt", "quelle", "fundstelle"},
    "oxford": {"oxford 3000", "oxford3000", "oxford", "oxford 5000"},
}


@dataclass
class Workbook:
    """Eingelesene Wortliste mit allen erkannten Units."""

    entries: list[VocabEntry] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    sheet: str = ""

    @property
    def units(self) -> list[int]:
        """Alle vorkommenden Unit-Nummern (0 = Starter Unit), aufsteigend."""
        return sorted({e.unit for e in self.entries if e.unit is not None})

    def unit_entries(self, unit: int) -> list[VocabEntry]:
        return [e for e in self.entries if e.unit == unit]

    def entries_before(self, unit: int) -> list[VocabEntry]:
        """Vokabular früherer Units - gilt als bereits gelernt."""
        return [e for e in self.entries if e.unit is not None and e.unit < unit]

    def unit_label(self, unit: int) -> str:
        return "Starter Unit" if unit == 0 else f"Unit {unit}"


def strip_page_reference(section: str) -> str:
    """Entfernt Seitenangaben, damit die Unit-Nummer nicht mit ihnen verwechselt wird."""
    s = str(section).replace("\n", " ").strip()
    previous = None
    while previous != s:
        previous = s
        s = _PAGE_SUFFIX.sub("", s).strip()
        s = _PAGE_PAREN.sub("", s).strip()
    return s.strip().strip(",;").strip()


def parse_unit(section: str) -> tuple[int | None, str]:
    """Bestimmt die Unit aus einer Sektionsangabe.

    Gibt ``(unit, bereinigtes_label)`` zurück; ``0`` steht für die Starter Unit,
    ``None`` wenn keine Unit erkennbar ist.
    """
    label = strip_page_reference(section)
    if not label:
        return None, ""
    if _STARTER.search(label):
        return 0, label
    m = _UNIT_NUM.search(label)
    if m:
        return int(m.group(1)), label
    m = _TRAILING_NUM.search(label)
    if m:
        return int(m.group(1)), label
    return None, label


def _norm(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _find_header(rows: Sequence[Sequence[object]]) -> tuple[int, dict[str, int]] | None:
    """Sucht die Kopfzeile und ordnet die benötigten Spalten zu."""
    for idx, row in enumerate(rows[:40]):
        cells = {c: _norm(v).lower().strip("*: ") for c, v in enumerate(row)}
        mapping: dict[str, int] = {}
        for name, aliases in _HEADER_ALIASES.items():
            for col, text in cells.items():
                if not text:
                    continue
                first_line = text.split("\n")[0].strip()
                if first_line in aliases or text in aliases:
                    mapping.setdefault(name, col)
                    break
        if "english" in mapping and "german" in mapping:
            return idx, mapping
    return None


def _looks_like_block_header(row: Sequence[object], cols: dict[str, int]) -> str | None:
    """Blocküberschrift = Text in der ersten Spalte, aber keine Vokabeldaten."""
    english = _norm(row[cols["english"]]) if cols["english"] < len(row) else ""
    german = _norm(row[cols["german"]]) if cols["german"] < len(row) else ""
    if english or german:
        return None
    for value in row:
        text = _norm(value)
        if text:
            return text
    return None


def read_workbook(source: str | Path | IO[bytes], sheet: str | None = None) -> Workbook:
    """Liest die Wortliste ein und ordnet jeden Eintrag einer Unit zu."""
    wb = load_workbook(source, data_only=True, read_only=True)
    try:
        sheets = [wb[sheet]] if sheet else list(wb.worksheets)
        for ws in sheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            found = _find_header(rows)
            if found is None:
                continue
            header_idx, cols = found
            return _parse_rows(rows, header_idx, cols, ws.title)
    finally:
        wb.close()

    raise ValueError(
        "In der Excel-Datei wurde keine Kopfzeile mit den Spalten 'English' und "
        "'Translation'/'Deutsch' gefunden. Bitte prüfen Sie das Tabellenblatt."
    )


def _parse_rows(
    rows: Sequence[Sequence[object]],
    header_idx: int,
    cols: dict[str, int],
    sheet_title: str,
) -> Workbook:
    book = Workbook(sheet=sheet_title)
    en_col, de_col = cols["english"], cols["german"]
    sec_col = cols.get("section")
    ox_col = cols.get("oxford")

    current_block: str | None = None
    block_unit: int | None = None
    seen: set[tuple[str, str, int | None]] = set()
    mismatches = 0

    for offset, raw in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        row = list(raw)
        while len(row) <= max(en_col, de_col, sec_col or 0, ox_col or 0):
            row.append(None)

        header_text = _looks_like_block_header(row, {"english": en_col, "german": de_col})
        if header_text is not None:
            current_block = header_text
            unit_from_block, _ = parse_unit(header_text)
            # "Culture"/"Projects" ohne Nummer setzen den Block-Kontext zurück.
            block_unit = unit_from_block
            continue

        english = _norm(row[en_col])
        german = _norm(row[de_col])
        if not english:
            continue
        if not german:
            book.warnings.append(
                f"Zeile {offset}: '{english}' hat keine deutsche Übersetzung und wird übersprungen."
            )
            continue

        section = _norm(row[sec_col]) if sec_col is not None else (current_block or "")
        unit, label = parse_unit(section) if section else (block_unit, current_block or "")

        # Zweite, unabhängige Quelle: die Blocküberschrift. Weichen beide ab,
        # gewinnt die Sektionsangabe (feiner), aber wir protokollieren es.
        if unit is None and block_unit is not None:
            unit = block_unit
            label = current_block or label
        elif unit is not None and block_unit is not None and unit != block_unit:
            mismatches += 1

        key = (english.lower(), german.lower(), unit)
        if key in seen:
            continue
        seen.add(key)

        book.entries.append(
            VocabEntry(
                english=english,
                german=german,
                section=section or (current_block or ""),
                unit=unit,
                block=current_block,
                oxford3000=bool(_norm(row[ox_col])) if ox_col is not None else False,
                row=offset,
                sheet=sheet_title,
            )
        )

    if mismatches:
        book.warnings.append(
            f"{mismatches} Einträge: Blocküberschrift und Spalte 'Section' nennen "
            "unterschiedliche Units. Es gilt die Angabe aus 'Section'."
        )
    unassigned = sum(1 for e in book.entries if e.unit is None)
    if unassigned:
        book.warnings.append(
            f"{unassigned} Einträge konnten keiner Unit zugeordnet werden und "
            "bleiben unberücksichtigt."
        )
    if not book.entries:
        raise ValueError("Die Excel-Datei enthält keine auswertbaren Vokabelzeilen.")
    return book


def collect_unit_pool(book: Workbook, unit: int) -> list[VocabEntry]:
    """Alle Einträge der gewünschten Unit - inklusive Culture/Project/Extra-Teile."""
    pool = book.unit_entries(unit)
    if not pool:
        available = ", ".join(book.unit_label(u) for u in book.units)
        raise ValueError(
            f"Für '{book.unit_label(unit)}' wurden keine Vokabeln gefunden. "
            f"Verfügbar sind: {available}"
        )
    return pool


def parse_unit_input(text: str) -> int:
    """Wandelt Nutzereingaben wie 'Unit 4', '4', 'unit4' oder 'Starter' in eine Nummer."""
    s = str(text).strip()
    if not s:
        raise ValueError("Bitte geben Sie eine Unit an, z. B. 'Unit 4'.")
    if _STARTER.search(s):
        return 0
    m = re.search(r"(\d{1,2})", s)
    if not m:
        raise ValueError(f"'{text}' konnte nicht als Unit gelesen werden. Beispiel: 'Unit 4'.")
    return int(m.group(1))


def iter_headwords(entries: Iterable[VocabEntry]) -> set[str]:
    return {e.english.strip().lower() for e in entries}
