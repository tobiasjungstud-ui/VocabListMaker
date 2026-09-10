"""Word-Ausgabe: Layout genau wie in der Vorlage."""

from __future__ import annotations

import re
import zipfile
from io import BytesIO

import pytest
from docx import Document

from vocablistmaker.config import Settings
from vocablistmaker.docx_writer import (
    COLUMN_WIDTHS,
    ROW_HEIGHT,
    TABLE_WIDTH,
    suggested_filename,
    to_bytes,
)
from vocablistmaker.models import TestItem, TestPair


@pytest.fixture
def pair() -> TestPair:
    result = TestPair(unit=4, unit_label="Unit 4")
    result.test1 = [
        TestItem(i, "Abfall", "rubbish", "We collect rubbish every Friday.", "rubbish")
        for i in range(1, 31)
    ]
    result.test2 = [
        TestItem(i, "überzeugen", "persuade", "She tried to persuade her parents.", "persuade")
        for i in range(1, 31)
    ]
    return result


@pytest.fixture
def document(pair: TestPair) -> Document:
    return Document(BytesIO(to_bytes(pair)))


def test_two_tables_with_header_and_thirty_rows(document: Document) -> None:
    assert len(document.tables) == 2
    for table in document.tables:
        assert len(table.rows) == 31  # Kopfzeile + 30 Einträge
        assert len(table.columns) == 4


def test_headings_match_template(document: Document) -> None:
    assert [p.text for p in document.paragraphs] == ["Test 1", "Test 2"]


def test_column_titles(document: Document) -> None:
    header = [c.text for c in document.tables[0].rows[0].cells]
    assert header == ["Nr.", "Deutsch", "English", "Example sentence"]


def test_numbering_is_continuous_per_test(document: Document) -> None:
    for table in document.tables:
        numbers = [row.cells[0].text for row in table.rows[1:]]
        assert numbers == [str(i) for i in range(1, 31)]


def test_row_numbers_are_bold_and_right_aligned(document: Document) -> None:
    cell = document.tables[0].rows[1].cells[0]
    run = cell.paragraphs[0].runs[0]
    assert run.bold is True
    assert cell.paragraphs[0].alignment is not None


def test_target_word_is_bold_inside_the_sentence(document: Document) -> None:
    runs = document.tables[0].rows[1].cells[3].paragraphs[0].runs
    bold_text = "".join(r.text for r in runs if r.bold)
    assert bold_text == "rubbish"
    assert "".join(r.text for r in runs) == "We collect rubbish every Friday."


def test_page_and_table_geometry_match_template(pair: TestPair) -> None:
    xml = zipfile.ZipFile(BytesIO(to_bytes(pair))).read("word/document.xml").decode("utf-8")
    assert '<w:pgSz w:w="11906" w:h="16838"/>' in xml
    assert 'w:top="728"' in xml and 'w:left="1417"' in xml
    assert f'<w:tblW w:w="{TABLE_WIDTH}" w:type="dxa"/>' in xml
    assert re.findall(r'<w:gridCol w:w="(\d+)"/>', xml)[:4] == [str(w) for w in COLUMN_WIDTHS]
    assert set(re.findall(r'<w:trHeight w:val="(\d+)"', xml)) == {str(ROW_HEIGHT)}


def test_font_is_century_gothic_at_eleven_point(pair: TestPair) -> None:
    xml = zipfile.ZipFile(BytesIO(to_bytes(pair))).read("word/document.xml").decode("utf-8")
    assert 'w:ascii="Century Gothic"' in xml
    assert '<w:sz w:val="22"/>' in xml  # 22 Halbpunkt = 11 pt


def test_only_horizontal_borders(pair: TestPair) -> None:
    xml = zipfile.ZipFile(BytesIO(to_bytes(pair))).read("word/document.xml").decode("utf-8")
    assert '<w:left w:val="nil"/>' in xml
    assert '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>' in xml


def test_settings_can_switch_off_bold_target(pair: TestPair) -> None:
    settings = Settings(bold_target_word=False)
    document = Document(BytesIO(to_bytes(pair, settings)))
    runs = document.tables[0].rows[1].cells[3].paragraphs[0].runs
    assert not any(r.bold for r in runs)


def test_suggested_filename(pair: TestPair) -> None:
    assert suggested_filename(pair) == "Vocabulary_Unit_4.docx"
